# -*- coding: utf-8 -*-
"""
Genera plantillas .docx web (docxtpl) para hm_divisa / BGT desde:
  - CONTRATO - EJEMPLO CONDUCTOR.rtf.doc
  - EJMEPLO CONTRATO SUPERVISOR.docx

Marcadores válidos: los de contratos_docs.MARCADORES_AYUDA
({{ trabajador }}, {{ dni }}, {{ empresa }}, {{ cargo }}, ...).
Además de los bookmarks legacy, se marcan los datos duros del
encabezado (empresa/RUC/domicilio/representante) que trae el SP
sp_pr_generarcontrato_datos_web. No se marca {{ nacionalidad }}:
ese SP la devuelve vacía.
"""
from __future__ import annotations

import re
import shutil
import sys
import zipfile
from pathlib import Path

import pythoncom
import win32com.client  # type: ignore
from docx import Document

ROOT = Path(r"C:\PROYECTOS\PLANILLAS")
sys.path.insert(0, str(ROOT))
from contratos_docs import MARCADORES_LEGACY_MAP, plantillas_dir  # noqa: E402

CIA = "BGT"
CONDUCTOR_SRC = ROOT / "AUXILIARES" / "CONTRATOS" / "CONTRATO - EJEMPLO CONDUCTOR.rtf.doc"
CONDUCTOR_OUT = ROOT / "AUXILIARES" / "CONTRATOS" / "CONTRATO_CONDUCTOR_DIVISA.docx"
SUPERVISOR_SRC = ROOT / "AUXILIARES" / "CONTRATOS" / "EJMEPLO CONTRATO SUPERVISOR.docx"
SUPERVISOR_OUT = ROOT / "AUXILIARES" / "CONTRATOS" / "CONTRATO_SUPERVISOR_DIVISA.docx"

# Bookmarks del conductor (todos ya están en MARCADORES_LEGACY_MAP; se
# dejan explícitos por si el mapa legacy cambia)
CONDUCTOR_EXTRA = {
    "FIRMACONTRATOTEXTO1": "fecha_firma",
    "PERFILCONTABLE": "centro_costo",
    "FECHAINICONTRATO": "inicio_contrato",
    "FECHAFINCONTRATO": "fin_contrato",
    "NOMBRETRAB": "trabajador",
    "NRODOCTRAB": "dni",
    "DIRECCIONTRAB1": "direccion",
    "CARGOTRAB1": "cargo",
    "SALARIONUMERO": "sueldo",
    "SALARIO": "sueldo_letras",
}

SUPERVISOR_BOOKMARKS = {
    "PERFILCONTABLE": "centro_costo",
    "CARGOTRAB1": "cargo",
    "SALARIONUMERO": "sueldo",
    "SALARIO": "sueldo_letras",
    "FIRMACONTRATOTEXTO1": "fecha_firma",
}

_SAMPLE_AFTER_MARKER = [
    (r"(\{\{\s*trabajador\s*\}\})\s*[A-ZÁÉÍÓÚÑa-záéíóúñ. ]+?(?=\s*,)", r"\1"),
    (r"(\{\{\s*dni\s*\}\})\s*\d+", r"\1"),
    (r"(\{\{\s*direccion\s*\}\})\s*[^,]+(?=\s*,)", r"\1"),
    (r"(\{\{\s*centro_costo\s*\}\})\s*\S+", r"\1"),
    # cargo puede ser varias palabras en MAYÚSCULAS (SUPERVISOR DE OPERACIONES)
    (r"(\{\{\s*cargo\s*\}\})\s*[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ ]*?(?=\s*,)", r"\1"),
    (r"(\{\{\s*cargo\s*\}\})\s*\S+", r"\1"),
    (
        r"(\{\{\s*inicio_contrato\s*\}\})\s*\d{1,2}\s+de\s+[A-Za-zÁÉÍÓÚáéíóúñÑ]+\s+del?\s+\d{2,4}",
        r"\1",
    ),
    (
        r"(\{\{\s*fin_contrato\s*\}\})\s*\d{1,2}\s+de\s+[A-Za-zÁÉÍÓÚáéíóúñÑ]+\s+del?\s+\d{2,4}",
        r"\1",
    ),
    (r"(\{\{\s*periodo_contrato\s*\}\})\s*(?:tres|\d+)", r"\1", re.IGNORECASE),
    (r"(\{\{\s*sueldo\s*\}\})\s*[\d,]+\.?\d*", r"\1"),
    (r"(\{\{\s*sueldo_letras\s*\}\})\s*\([^)]*\)", r"\1"),
    (r"(\{\{\s*sueldo_letras\s*\}\})\s*[A-ZÁÉÍÓÚÑa-záéíóúñ. ]+con\s+\d{2}/100", r"\1", re.IGNORECASE),
    (
        r"(\{\{\s*fecha_firma\s*\}\})\s*\d{1,2}\s+de\s+[A-Za-zÁÉÍÓÚáéíóúñÑ]+\s+del?\s+\d{4}",
        r"\1",
    ),
]


def _wd_docx():
    return 12


def _marker(web: str) -> str:
    return "{{ " + web + " }}"


def _set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        for r in paragraph.runs:
            r.text = ""
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def _collect_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    return "\n".join(parts)


def _strip_sample_after_markers(text: str) -> str:
    out = text
    for item in _SAMPLE_AFTER_MARKER:
        pat, repl = item[0], item[1]
        flags = item[2] if len(item) > 2 else 0
        out = re.sub(pat, repl, out, flags=flags)
    return out


def enrich_docx(path: Path, legacy_names: list[str]) -> list[str]:
    doc = Document(str(path))
    for p in doc.paragraphs:
        t = _strip_sample_after_markers(p.text)
        for leg in legacy_names:
            t = re.sub(rf"\b{re.escape(leg)}\b\s*", "", t)
        t = re.sub(r"[ \t]{2,}", " ", t)
        if t != p.text:
            _set_paragraph_text(p, t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = _strip_sample_after_markers(p.text)
                    for leg in legacy_names:
                        t = re.sub(rf"\b{re.escape(leg)}\b\s*", "", t)
                    t = re.sub(r"[ \t]{2,}", " ", t)
                    if t != p.text:
                        _set_paragraph_text(p, t)
    doc.save(str(path))
    return sorted(set(re.findall(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}", _collect_text(doc))))


def _word_app():
    pythoncom.CoInitialize()
    # EnsureDispatch (tipado): con Dispatch dinámico, Word 12 pierde
    # Documents.Open tras el primer SaveAs2/Close
    try:
        word = win32com.client.gencache.EnsureDispatch("Word.Application")
    except Exception:
        word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    return word


def _word_quit(word) -> None:
    if word is not None:
        try:
            word.Quit()
        except Exception:
            pass
    pythoncom.CoUninitialize()


def convert_with_word_bookmarks(src: Path, dest: Path, mapping: dict[str, str]) -> int:
    """Convierte .doc (RTF) a .docx con Word (una sola apertura) y luego
    reemplaza los bookmarks en el XML, sin reabrir Word: Word 12 se cae
    (RPC) al abrir un segundo documento en la misma sesión COM."""
    word = None
    doc = None
    try:
        word = _word_app()
        doc = word.Documents.Open(
            FileName=str(src.resolve()),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
        )
        shapes = doc.Shapes.Count
        doc.SaveAs2(str(dest.resolve()), FileFormat=_wd_docx())
        doc.Close(False)
        doc = None
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass
        _word_quit(word)
    full_map = {**MARCADORES_LEGACY_MAP, **mapping}
    convert_docx_xml_bookmarks(dest, full_map)
    return shapes


def convert_docx_xml_bookmarks(path: Path, mapping: dict[str, str]) -> None:
    """Inserta {{ marcador }} entre bookmarkStart/bookmarkEnd en document.xml."""
    with zipfile.ZipFile(path, "r") as zin:
        payload = {name: zin.read(name) for name in zin.namelist()}
    xml = payload["word/document.xml"].decode("utf-8")
    for bk_name, web in mapping.items():
        marker = _marker(web)
        pat = (
            rf"(<w:bookmarkStart[^>]+w:name=\"{re.escape(bk_name)}\"[^>]*/>)"
            rf"(.*?)"
            rf"(<w:bookmarkEnd[^>]+w:id=\"(\d+)\"[^>]*/>)"
        )

        def repl(m, marker=marker):
            return f"{m.group(1)}<w:r><w:t xml:space=\"preserve\">{marker}</w:t></w:r>{m.group(3)}"

        xml2, n = re.subn(pat, repl, xml, count=1, flags=re.DOTALL)
        if n:
            xml = xml2
            print(f"  OK XML {bk_name} -> {marker}")
    payload["word/document.xml"] = xml.encode("utf-8")
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in payload.items():
            zout.writestr(name, data)


def convert_supervisor_docx_xml(src: Path, dest: Path) -> None:
    """Reemplaza contenido entre bookmarkStart/bookmarkEnd en document.xml."""
    shutil.copy2(src, dest)
    full_map = {**MARCADORES_LEGACY_MAP, **SUPERVISOR_BOOKMARKS}
    convert_docx_xml_bookmarks(dest, full_map)


# Encabezado duro (empresa/representante de VIAMERICA) común a ambos
# contratos: pasa a marcadores del SP de datos. doc_representante ya
# trae el tipo + número ("DNI N° 08777726"), por eso va sin "DNI N°".
_EMPRESA_RE = (
    r"VIAMERICA SAC,\s*identificada con RUC N[°º]\s*20425718577,\s*"
    r"(?:con domicilio para estos efectos en|y domicilio fiscal en)\s*"
    r"[^,]+?(?:Distrito de [^,]+,\s*Provincia de [^,]+,\s*Departamento de [^,]+)?,?\s*"
    r"debidamente representada por su Gerente General,\s*Sr\.\s*"
    r"BUSTAMANTE MAZZINI JOSE ERNESTO FERNANDO,\s*identificado con DNI N[°º]\s*08777726"
)
_EMPRESA_REPL = (
    "{{ empresa }}, identificada con RUC N° {{ ruc }}, con domicilio en "
    "{{ domicilio_empresa }}, debidamente representada por su "
    "{{ cargo_representante }}, Sr. {{ representante }}, identificado con "
    "{{ doc_representante }}"
)


def _apply_patches(path: Path, replacements: list[tuple[str, str]]) -> None:
    doc = Document(str(path))
    for p in doc.paragraphs:
        t = p.text
        orig = t
        for pat, repl in replacements:
            t = re.sub(pat, repl, t, flags=re.IGNORECASE)
        if t != orig:
            _set_paragraph_text(p, t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = p.text
                    orig = t
                    for pat, repl in replacements:
                        t = re.sub(pat, repl, t, flags=re.IGNORECASE)
                    if t != orig:
                        _set_paragraph_text(p, t)
    doc.save(str(path))


def patch_common(path: Path) -> None:
    """Marcadores sobre datos duros compartidos por conductor/supervisor."""
    _apply_patches(
        path,
        [
            (_EMPRESA_RE, _EMPRESA_REPL),
            # el tipo de documento también sale del trabajador (DNI/CE/…)
            (
                r"identificado con DNI N[°º]\s*(\{\{\s*dni\s*\}\})",
                r"identificado con {{ tipo_documento }} N° \1",
            ),
            # la limpieza del ejemplo se come la coma tras el cargo
            (r"(\{\{\s*cargo\s*\}\})\s+contratación", r"\1, contratación"),
            # sueldo_letras ya incluye "CON 00/100 SOLES"
            (r"(\{\{\s*sueldo_letras\s*\}\})\s+Soles\b", r"\1"),
            # S/. {{ sueldo }} {{ sueldo_letras }} -> S/. {{ sueldo }} ({{ sueldo_letras }})
            (
                r"(S/\.\s*\{\{\s*sueldo\s*\}\})\s+(?=\{\{\s*sueldo_letras)",
                r"\1 (",
            ),
        ],
    )
    # cerrar el paréntesis abierto arriba solo en el párrafo del sueldo
    doc = Document(str(path))
    for p in doc.paragraphs:
        if "{{ sueldo }} ({{ sueldo_letras }}" in p.text and "{{ sueldo_letras }})" not in p.text:
            _set_paragraph_text(
                p, p.text.replace("{{ sueldo_letras }}", "{{ sueldo_letras }})", 1)
            )
    doc.save(str(path))


def patch_conductor_hardcoded(path: Path) -> None:
    _apply_patches(
        path,
        [
            # PRIMERO: "profesión/ oficio CONDUCTOR" (sin bookmark)
            (r"(oficio\s+)CONDUCTOR\b", r"\1{{ cargo }}"),
        ],
    )


def patch_supervisor_hardcoded(path: Path) -> None:
    doc = Document(str(path))
    replacements = [
        (
            r"del Sr\.\s*[^,]+,\s*identificado con DNI N[°º]\s*\d+,\s*domiciliado en\s*[^,]+",
            "del Sr. {{ trabajador }}, identificado con DNI N° {{ dni }}, domiciliado en {{ direccion }}",
        ),
        (
            r"plazo de duraci[oó]n de\s+[^,]+,\s*el mismo que regir[aá]\s+a partir del\s+\d{1,2}\s+de\s+\w+\s+del\s+\d{2,4},\s*fecha en que EL TRABAJADOR debe empezar sus labores hasta el\s+\d{1,2}\s+de\s+\w+\s+del\s+\d{4}",
            "plazo de duración de {{ periodo_contrato }}, el mismo que regirá a partir del {{ inicio_contrato }}, fecha en que EL TRABAJADOR debe empezar sus labores hasta el {{ fin_contrato }}",
        ),
        (
            r"en Lima el\s+\d{1,2}\s+de\s+\w+\s+del\s+\d{4}",
            "en Lima el {{ fecha_firma }}",
        ),
    ]
    for p in doc.paragraphs:
        t = p.text
        orig = t
        for pat, repl in replacements:
            t = re.sub(pat, repl, t, flags=re.IGNORECASE)
        if t != orig:
            _set_paragraph_text(p, t)
    doc.save(str(path))


def install(path: Path, filename: str) -> Path:
    dest = plantillas_dir(CIA, database="hm_divisa") / filename
    shutil.copy2(path, dest)
    print("Instalada:", dest)
    return dest


def _final_markers(path: Path) -> list[str]:
    doc = Document(str(path))
    return sorted(set(re.findall(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}", _collect_text(doc))))


def build_conductor() -> list[str]:
    if not CONDUCTOR_SRC.is_file():
        raise SystemExit(f"No existe: {CONDUCTOR_SRC}")
    print("\n=== CONDUCTOR ===")
    shapes = convert_with_word_bookmarks(CONDUCTOR_SRC, CONDUCTOR_OUT, CONDUCTOR_EXTRA)
    print(f"Shapes (firma/logo): {shapes}")
    legacy = list({*MARCADORES_LEGACY_MAP.keys(), *CONDUCTOR_EXTRA.keys()})
    markers = enrich_docx(CONDUCTOR_OUT, legacy)
    print("Marcadores tras bookmarks:", markers)
    patch_conductor_hardcoded(CONDUCTOR_OUT)
    patch_common(CONDUCTOR_OUT)
    markers = _final_markers(CONDUCTOR_OUT)
    print("Marcadores finales:", markers)
    install(CONDUCTOR_OUT, "CONTRATO_CONDUCTOR_DIVISA.docx")
    return markers


def build_supervisor() -> list[str]:
    if not SUPERVISOR_SRC.is_file():
        raise SystemExit(f"No existe: {SUPERVISOR_SRC}")
    print("\n=== SUPERVISOR ===")
    convert_supervisor_docx_xml(SUPERVISOR_SRC, SUPERVISOR_OUT)
    patch_supervisor_hardcoded(SUPERVISOR_OUT)
    legacy = list({*MARCADORES_LEGACY_MAP.keys(), *SUPERVISOR_BOOKMARKS.keys()})
    markers = enrich_docx(SUPERVISOR_OUT, legacy)
    print("Marcadores tras bookmarks:", markers)
    patch_common(SUPERVISOR_OUT)
    markers = _final_markers(SUPERVISOR_OUT)
    print("Marcadores finales:", markers)
    install(SUPERVISOR_OUT, "CONTRATO_SUPERVISOR_DIVISA.docx")
    return markers


def main() -> None:
    build_conductor()
    build_supervisor()
    print("\nLISTO")


if __name__ == "__main__":
    main()
