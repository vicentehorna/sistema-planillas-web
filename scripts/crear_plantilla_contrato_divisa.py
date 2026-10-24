# -*- coding: utf-8 -*-
"""
Convierte el contrato ejemplo Divisa (RTF/.doc) a plantilla .docx web
con marcadores Jinja (docxtpl) definidos en contratos_docs.MARCADORES_AYUDA.
"""
from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

import win32com.client  # type: ignore
from docx import Document

ROOT = Path(r"C:\PROYECTOS\PLANILLAS")
sys.path.insert(0, str(ROOT))
from contratos_docs import MARCADORES_LEGACY_MAP, plantillas_dir  # noqa: E402

SRC = ROOT / "AUXILIARES" / "CONTRATOS" / "CONTRATO - EJEMPLO AUXILIAR.doc"
OUT_AUX = ROOT / "AUXILIARES" / "CONTRATOS" / "CONTRATO_EJEMPLO_AUXILIAR_DIVISA.docx"
CIA = "BGT"  # VIAMERICA SAC en hm_divisa

# Bookmarks del DOC que no están en el mapa legacy estándar
EXTRA_LEGACY = {
    "FIRMACONTRATOTEXTO1": "fecha_firma",
    "PERFILCONTABLE": "centro_costo",
    "FECHAINICONTRATO1": "inicio_contrato",
    "NOMBRETRAB1": "trabajador",
    "NRODOCTRAB1": "dni",
    "CARGOTRAB": "cargo",
    "CARGOTRAB2": "cargo",
    "CARGOTRAB3": "cargo",
    "PERIODOCONTRATO": "periodo_contrato",
}


def _wd_format_docx():
    # 12 = wdFormatXMLDocument (.docx)
    return 12


def convert_and_replace_bookmarks() -> Path:
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    tmp_docx = ROOT / "AUXILIARES" / "CONTRATOS" / "_tmp_divisa_auxiliar.docx"
    try:
        doc = word.Documents.Open(
            FileName=str(SRC),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
        )
        # Guardar primero como docx para trabajar limpio
        doc.SaveAs2(FileName=str(tmp_docx), FileFormat=_wd_format_docx())
        doc.Close(False)
        doc = word.Documents.Open(
            FileName=str(tmp_docx),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
        )

        mapping = {**MARCADORES_LEGACY_MAP, **EXTRA_LEGACY}
        names = [b.Name for b in doc.Bookmarks]
        print("Bookmarks encontrados:", names)
        for bk_name in list(names):
            if str(bk_name).startswith("_"):
                continue
            web = mapping.get(str(bk_name).upper())
            if not web:
                print(f"  SKIP sin mapa: {bk_name}")
                continue
            marker = "{{ " + web + " }}"
            try:
                rng = doc.Bookmarks(bk_name).Range
                rng.Text = marker
                # Re-crear bookmark sobre el marcador (opcional; no necesario para docxtpl)
                print(f"  OK {bk_name} -> {marker}")
            except Exception as ex:
                print(f"  ERR {bk_name}: {ex}")

        doc.SaveAs2(FileName=str(OUT_AUX), FileFormat=_wd_format_docx())
        doc.Close(False)
    finally:
        word.Quit()
        if tmp_docx.exists():
            try:
                tmp_docx.unlink()
            except Exception:
                pass
    return OUT_AUX


def _set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        for r in paragraph.runs:
            r.text = ""
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


# Tras reemplazar bookmarks de longitud 0 / parcial, suele quedar el valor de ejemplo.
_SAMPLE_AFTER_MARKER = [
    (r"(\{\{\s*trabajador\s*\}\})\s*[A-ZÁÉÍÓÚÑa-záéíóúñ. ]+?(?=\s*,)", r"\1"),
    (r"(\{\{\s*dni\s*\}\})\s*\d+", r"\1"),
    (r"(\{\{\s*direccion\s*\}\})\s*[^,]+(?=\s*,)", r"\1"),
    (r"(\{\{\s*centro_costo\s*\}\})\s*\S+", r"\1"),
    (r"(\{\{\s*cargo\s*\}\})\s*\S+", r"\1"),
    (
        r"(\{\{\s*inicio_contrato\s*\}\})\s*\d{1,2}\s+de\s+[A-Za-zÁÉÍÓÚáéíóúñÑ]+\s+del?\s+\d{4}",
        r"\1",
    ),
    (
        r"(\{\{\s*fin_contrato\s*\}\})\s*\d{1,2}\s+de\s+[A-Za-zÁÉÍÓÚáéíóúñÑ]+\s+del?\s+\d{4}\s*",
        r"\1",
    ),
    (r"(\{\{\s*sueldo\s*\}\})\s*[\d,]+\.?\d*", r"\1"),
    (r"(\{\{\s*sueldo_letras\s*\}\})\s*\([^)]*\)", r"\1"),
    (
        r"(\{\{\s*fecha_firma\s*\}\})\s*\d{1,2}\s+de\s+[A-Za-zÁÉÍÓÚáéíóúñÑ]+\s+del?\s+\d{4}",
        r"\1",
    ),
]


def _strip_sample_after_markers(text: str) -> str:
    out = text
    for pat, repl in _SAMPLE_AFTER_MARKER:
        out = re.sub(pat, repl, out, flags=re.IGNORECASE)
    return out


def enrich_with_web_markers(path: Path) -> None:
    """Limpia restos de bookmarks/valores de ejemplo; deja plantilla lista para docxtpl."""
    doc = Document(str(path))
    joined = "\n".join(p.text for p in doc.paragraphs)
    present = set(re.findall(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}", joined))
    print("Marcadores en plantilla:", sorted(present))

    # Limpiar nombres de bookmark residuales + valores de ejemplo pegados tras {{ marcador }}
    legacy_names = list({*MARCADORES_LEGACY_MAP.keys(), *EXTRA_LEGACY.keys()})
    for p in doc.paragraphs:
        t = p.text
        orig = t
        for leg in legacy_names:
            t = re.sub(rf"\b{re.escape(leg)}\b\s*", "", t)
        t = _strip_sample_after_markers(t)
        t = re.sub(r"[ \t]{2,}", " ", t)
        if t != orig:
            _set_paragraph_text(p, t)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    t = p.text
                    orig = t
                    for leg in legacy_names:
                        t = re.sub(rf"\b{re.escape(leg)}\b\s*", "", t)
                    t = _strip_sample_after_markers(t)
                    t = re.sub(r"[ \t]{2,}", " ", t)
                    if t != orig:
                        _set_paragraph_text(p, t)

    doc.save(str(path))
    joined2 = "\n".join(p.text for p in Document(str(path)).paragraphs)
    present2 = sorted(set(re.findall(r"\{\{\s*([a-zA-Z0-9_]+)\s*\}\}", joined2)))
    print("Marcadores finales:", present2)


def install_for_cia(path: Path, cia: str) -> Path:
    dest_dir = plantillas_dir(cia, database="hm_divisa")
    dest = dest_dir / "CONTRATO_EJEMPLO_AUXILIAR_DIVISA.docx"
    shutil.copy2(path, dest)
    print("Instalada en", dest)
    return dest


def main():
    if not SRC.is_file():
        raise SystemExit(f"No existe: {SRC}")
    out = convert_and_replace_bookmarks()
    enrich_with_web_markers(out)
    install_for_cia(out, CIA)
    print("LISTO:", out)


if __name__ == "__main__":
    main()
